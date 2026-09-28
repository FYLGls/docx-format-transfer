# -*- coding: utf-8 -*-
"""parse_requirements.py — 从"格式要求"文档/图片/粘贴文字解析结构化格式规则

用法：
  python parse_requirements.py <要求.docx|.pdf|.txt|.md|.png|.jpg> [-o req.json]
                               [--merge profile.json [-o merged.json]]       # 文件模式
  python parse_requirements.py --text "正文宋体小四，1.5倍行距"                 # 直接解析文字

提取文本：docx/txt/md 直接读；pdf 先文字层（pypdf/PyMuPDF），扫描件渲染后 OCR；
图片用 rapidocr-onnxruntime（未安装时提示安装或改用粘贴文字）。
解析结果 rules 以 profile 路径为键，value 优先级：用户修改 > 格式要求 > 模板实测。
"""
from __future__ import annotations

import argparse
import io
import json
import os
import re
import sys

SOURCE = "格式要求"

HAO_TO_PT = {"初号": 42, "小初": 36, "一号": 26, "小一": 24, "二号": 22, "小二": 18,
             "三号": 16, "小三": 15, "四号": 14, "小四": 12, "五号": 10.5,
             "小五": 9, "六号": 7.5, "小六": 6.5}
CN_FONTS = ["宋体", "黑体", "楷体", "楷体_GB2312", "仿宋", "仿宋_GB2312", "微软雅黑",
            "Times New Roman", "Arial", "Calibri", "Consolas", "Courier New"]
ROLE_MAP = [  # (别名正则, profile 前缀组)；泛指"标题"特判
    (r"一级标题|章标题|chapter", ["headings.level1"]),
    (r"二级标题|节标题", ["headings.level2"]),
    (r"三级标题", ["headings.level3"]),
    (r"正文", ["body"]),
    (r"图表题注|题注|图注|表注", ["captions.figure", "captions.table"]),
    (r"图题注|图注|图的?(标题|名称)", ["captions.figure"]),
    (r"表题注|表注|表的?(标题|名称)", ["captions.table"]),
    (r"表内|表格内|表格文字", ["tables"]),
    (r"页码", ["page_numbers.body"]),
    (r"参考文献", ["bibliography"]),
    (r"目录", ["toc"]),
    (r"代码", ["code_blocks"]),
]


def _role_prefixes(sentence: str, levels_in_use=None) -> list[str]:
    if re.search(r"标题", sentence) and not re.search(r"一级|二级|三级|章标题|节标题", sentence):
        # 泛指"标题"→ 模板在用的所有级别（levels_in_use 由调用方传入）
        lv = levels_in_use or [1, 2, 3]
        return [f"headings.level{x}" for x in lv]
    for pat, prefixes in ROLE_MAP:
        if re.search(pat, sentence):
            return prefixes
    return []


def _num(sentence: str, pat: str, cast=float):
    m = re.search(pat, sentence)
    if not m:
        return None
    try:
        return cast(m.group(1))
    except (ValueError, IndexError):
        return None


class RequirementParser:
    """把中文格式要求句子解析为 {profile路径: 规则}。levels_in_use 用于泛指标题。"""

    def __init__(self, levels_in_use=None):
        self.levels_in_use = levels_in_use

    def parse(self, text: str) -> dict:
        """句子 → [结构族整句匹配] + [角色属性子句匹配(缺角色子句继承整句角色)]。"""
        rules: dict = {}
        unparsed: list = []
        for raw in re.split(r"[。\n]+", text or ""):
            s = raw.strip()
            if not s:
                continue
            sent_hits: dict = {}
            for path, node in self._structural(s):
                sent_hits[path] = node
            sent_roles = self._roles_of(s)
            for clause in re.split(r"[，,、]+", s):
                c = clause.strip()
                if not c:
                    continue
                roles = self._roles_of(c) or sent_roles
                for path, node in self._role_attrs(c, roles):
                    sent_hits[path] = node
            if sent_hits:
                for path, node in sent_hits.items():
                    node["evidence"] = s
                    rules[path] = node
            elif re.search(r"格式|字体|字号|页边距|行距|缩进|页码|题注|三线|页眉|页脚|目录|排版",
                           s):
                unparsed.append(s)  # 像格式要求但没解析出规则 → 透明展示
        return {"rules": rules, "unparsed": unparsed}

    def _roles_of(self, s: str) -> list[str]:
        return _role_prefixes(s, self.levels_in_use)

    # ---------------------------------------------------------------- 结构族（整句）
    def _structural(self, s: str) -> list[tuple[str, dict]]:
        out: list[tuple[str, dict]] = []

        def add(path, value, what):
            out.append((path, {"value": value, "source": SOURCE, "what": what}))

        # 纸张/方向 + 页边距（两个匹配器并行累加，同句不互斥）
        size_hit = False
        if re.search(r"A4|A3|纸张|横向|纵向", s):
            if "A4" in s:
                add("page.size.width_cm", 21.0, "A4 纸")
                add("page.size.height_cm", 29.7, "A4 纸")
            if "横向" in s:
                add("page.size.orientation", "landscape", "横向")
            size_hit = True
        if "页边距" in s or "边距" in s:
            m_all = _num(s, r"均[为是]?\s*(\d+(?:\.\d+)?)\s*(?:cm|厘米|㎝)?")
            pairs = [(r"上(?:边距|侧)?[^0-9]{0,4}(\d+(?:\.\d+)?)", "top_cm"),
                     (r"下(?:边距|侧)?[^0-9]{0,4}(\d+(?:\.\d+)?)", "bottom_cm"),
                     (r"左(?:边距|侧)?[^0-9]{0,4}(\d+(?:\.\d+)?)", "left_cm"),
                     (r"右(?:边距|侧)?[^0-9]{0,4}(\d+(?:\.\d+)?)", "right_cm")]
            got = False
            for pat, key in pairs:
                v = _num(s, pat)
                if v is None:
                    v = m_all
                if v is not None:
                    add(f"page.margins.{key}", v, f"页边距 {key[:-3]}={v}cm")
                    got = True
            if got:
                size_hit = True
        if size_hit:
            return out
        # 行距 / 缩进（两个匹配器独立累加，互不吞并）
        m = re.search(r"(\d+(?:\.\d+)?)\s*倍\s*(?:的)?行距|行距[^0-9]{0,4}(\d+(?:\.\d+)?)\s*倍",
                      s)
        if m:
            v = float(m.group(1) or m.group(2))
            add("body.line_spacing", {"rule": "auto", "value": v}, f"{v} 倍行距")
        m = re.search(r"固定值\s*(\d+(?:\.\d+)?)\s*磅|行距[^0-9]{0,6}(\d+(?:\.\d+)?)\s*磅", s)
        if m and not any(p == "body.line_spacing" for p, _ in out):
            v = float(m.group(1) or m.group(2))
            add("body.line_spacing", {"rule": "exact", "value": v}, f"固定值 {v} 磅")
        m = re.search(r"首行缩进\s*(\d+(?:\.\d+)?)\s*字符", s)
        if m:
            add("body.first_line_indent_chars", float(m.group(1)),
                f"首行缩进 {m.group(1)} 字符")
        if any(p.startswith("body.line_spacing") or p == "body.first_line_indent_chars"
               for p, _ in out):
            return out
        # 三线表 / 表头（独立累加，不互斥）
        if "三线表" in s:
            add("tables.border_rule", "three-line", "三线表")
            add("tables.three_line.top_bottom_pt", 1.5, "三线表")
            add("tables.three_line.header_underline_pt", 0.75, "三线表")
        if re.search(r"重复表头|续表头|表头[^。]{0,6}重复", s):
            add("tables.repeat_header", True, "跨页重复表头")
        if any(p.startswith("tables.") for p, _ in out):
            return out
        # 页码
        if "页码" in s:
            hit = False
            if "居中" in s:
                add("page_numbers.body.position", "footer-center", "页码居中")
                hit = True
            elif re.search("居右|右侧", s):
                add("page_numbers.body.position", "footer-right", "页码居右")
                hit = True
            elif "外侧" in s:
                add("page_numbers.body.position", "footer-outside", "页码外侧")
                hit = True
            if re.search("阿拉伯", s):
                add("page_numbers.body.format", "decimal", "阿拉伯页码")
                hit = True
            if re.search("罗马", s):
                add("page_numbers.body.format", "lowerRoman", "罗马页码")
                hit = True
            if hit:
                return out
        # 目录
        m = re.search(r"([一二三1-3])\s*级目录", s)
        if m:
            depth = {"一": 1, "二": 2, "三": 3}.get(m.group(1)) or int(m.group(1))
            add("toc.depth", depth, f"{depth} 级目录")
            add("toc.mode", "field", "目录")
            return out
        # 代码
        if "代码" in s:
            hit = False
            if re.search("等宽|Consolas|Courier", s, re.I):
                add("code_blocks.font_ascii", "Consolas", "代码等宽字体")
                hit = True
            if re.search("装框|代码框|灰底|底纹", s):
                add("code_blocks.box_style", "single-cell-table", "代码装框")
                hit = True
            if hit:
                return out
        # 参考文献
        if "参考文献" in s and re.search(r"GB/?T\s*7714|\[\d+\]", s):
            add("bibliography.numbering_format", "[{n}]", "GB/T 7714 顺序编码")
            return out
        # 题注位置 / 编号（两个匹配器独立，不再互斥）
        if re.search(r"图(?:题注|注|的标题|标题)[^。]{0,10}?(下方|下面|之下|图后|图右)", s):
            add("captions.figure.position", "below", "图注在图下")
        elif re.search(r"图(?:题注|注|的标题|标题)[^。]{0,10}?(上方|上面|之上|图前|图左)", s):
            add("captions.figure.position", "above", "图注在图上")
        if re.search(r"表(?:题注|注|的标题|标题)[^。]{0,10}?(上方|上面|之上|表前|之前|顶部)", s):
            add("captions.table.position", "above", "表注在表上")
        elif re.search(r"表(?:题注|注|的标题|标题)[^。]{0,10}?(下方|下面|之下|表后|之后|底部)", s):
            add("captions.table.position", "below", "表注在表下")
        if re.search(r"[图表]\s*\d+\s*[-–]\s*\d+", s) and re.search(r"编号|图号|表号|格式", s):
            add("captions.figure.numbering.chapter_linked", True, "按章编号")
            add("captions.table.numbering.chapter_linked", True, "按章编号")
        # 每章另起页（无角色词时默认作用于一级标题）
        if re.search("每章另起|另起一页", s):
            add("headings.level1.page_break_before", True, "每章另起页")
            return out
        if out:
            return out
        # 题注位置族没命中时不再继续（字体族走子句通道）
        return []

    # ---------------------------------------------------------------- 角色属性（子句）
    def _role_attrs(self, c: str, roles: list[str]) -> list[tuple[str, dict]]:
        out: list[tuple[str, dict]] = []
        if not roles:
            return out

        def add(path, value, what):
            out.append((path, {"value": value, "source": SOURCE, "what": what}))
        for f in CN_FONTS:
            if f in c:
                for p in roles:
                    if p.startswith("headings") or p == "body":
                        add(f"{p}.font_eastasia", f, f"字体 {f}")
                    elif p.startswith("tables"):
                        add(f"{p}.cell_font_eastasia", f, f"表内字体 {f}")
                    elif p.startswith("captions"):
                        add(f"{p}.font_eastasia", f, f"字体 {f}")
                break
        hao = re.search(r"初号|小初|[一二三四五六七八]号|小[一二三四五六]", c)
        if hao:
            pt = HAO_TO_PT.get(hao.group(0))
            if pt:
                for p in roles:
                    if p.startswith("headings") or p == "body":
                        add(f"{p}.size_pt", pt, f"字号 {hao.group(0)}")
                    elif p.startswith("captions"):
                        add(f"{p}.size_pt", pt, f"字号 {hao.group(0)}")
                    elif p == "tables":
                        add(f"{p}.cell_size_pt", pt, f"表内字号 {hao.group(0)}")
        m = re.search(r"(\d+(?:\.\d+)?)\s*(?:磅|pt)\s*字|字[号形]{0,2}(\d+(?:\.\d+)?)\s*(?:磅|pt)",
                      c)
        if m:
            pt = float(m.group(1) or m.group(2))
            for p in roles:
                add(f"{p}.size_pt", pt, f"字号 {pt} 磅")
        if re.search("加粗|粗体", c):
            for p in roles:
                if p.startswith("headings") or p.startswith("captions"):
                    add(f"{p}.bold", True, "加粗")
        if "居中" in c:
            for p in roles:
                if p.startswith("headings") or p.startswith("captions"):
                    add(f"{p}.align", "center", "居中")
                elif p == "tables":
                    add("tables.cell_align", "center", "表内居中")
        elif "两端对齐" in c or "两端" in c:
            for p in roles:
                if p == "body":
                    add(f"{p}.align", "both", "两端对齐")
        if re.search("每章另起|另起一页", c):
            for p in roles:
                if p == "headings.level1":
                    add(f"{p}.page_break_before", True, "每章另起页")
        return out


# ---------------------------------------------------------------- 文本提取
def extract_text(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        return _text_docx(path)
    if ext in (".txt", ".md"):
        return open(path, encoding="utf-8", errors="replace").read()
    if ext == ".pdf":
        return _text_pdf(path)
    if ext in (".png", ".jpg", ".jpeg", ".bmp", ".webp"):
        return _text_ocr_image(path)
    raise ValueError(f"不支持的格式要求文件类型：{ext}（支持 docx/pdf/txt/md/png/jpg）")


def _text_docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _text_pdf(path: str) -> str:
    text = ""
    try:
        import pypdf
        rd = pypdf.PdfReader(path)
        text = "\n".join((pg.extract_text() or "") for pg in rd.pages)
    except ImportError:
        try:
            import fitz
            doc = fitz.open(path)
            text = "\n".join(pg.get_text() for pg in doc)
        except ImportError:
            raise ValueError("PDF 解析需要 pypdf：pip install pypdf")
    if len(text.strip()) < 20:  # 无文字层 → 扫描件走 OCR
        try:
            import fitz
            doc = fitz.open(path)
            ocr = _ocr_fn()
            pages = []
            for pg in doc:
                pix = pg.get_pixmap(dpi=200)
                img_path = path + f"_pg{pg.number}.png"
                pix.save(img_path)
                try:
                    pages.append(ocr(img_path))
                finally:
                    os.remove(img_path)
            text = "\n".join(pages)
        except ImportError:
            pass  # 没有 fitz 就只能用已有文字层
    return text


def _ocr_fn():
    """返回 ocr(image_path)->str；不可用则抛出带安装指引的 ValueError。"""
    try:
        from rapidocr_onnxruntime import RapidOCR
        engine = RapidOCR()

        def run(img_path):
            result, _ = engine(img_path)
            return "\n".join(x[1] for x in (result or []))
        return run
    except ImportError:
        raise ValueError("图片识别需要 OCR：pip install rapidocr-onnxruntime，"
                         "或直接把要求文字粘贴到输入框")


def _text_ocr_image(path: str) -> str:
    return _ocr_fn()(path)


# ---------------------------------------------------------------- 合并
def _get_value(profile: dict, path: str):
    cur = profile
    for p in path.split("."):
        if not isinstance(cur, dict):
            return None  # 中间节点缺失（如无页脚模板的 page_numbers.body=None）
        cur = cur.get(p)
    if isinstance(cur, dict) and "value" in cur and "rule" not in cur:
        return cur["value"]
    return cur


def _set_rule(profile: dict, path: str, node: dict):
    parts = path.split(".")
    cur = profile
    for p in parts[:-1]:
        cur = cur.setdefault(p, {})
    leaf = parts[-1]
    existing = cur.get(leaf)
    meta = {"source": SOURCE, "evidence": node["evidence"]}
    is_wrapper = isinstance(existing, dict) and "value" in existing and \
        set(existing) <= {"value", "source", "evidence", "what", "confidence"}
    if is_wrapper:
        existing.update({"value": node["value"], **meta})
    elif isinstance(existing, dict) and isinstance(node["value"], dict):
        # 字典值（如 line_spacing）：把键并进档案原 dict，附加来源元数据
        base = dict(existing)
        base.update(node["value"])
        base.update(meta)
        cur[leaf] = base
    else:
        cur[leaf] = {"value": node["value"], **meta}


def merge_profile(profile: dict, parsed: dict) -> dict:
    """把解析出的要求规则合并进模板档案：要求覆盖模板，冲突记 flag。"""
    merged = json.loads(json.dumps(profile, ensure_ascii=False))  # 深拷贝
    flags = merged.setdefault("flags", [])
    conflicts = []
    for path, node in parsed["rules"].items():
        old = _get_value(merged, path)
        new = node["value"]
        if old is not None and _differs(old, new):
            conflicts.append({"path": path, "template": old, "requirement": new,
                              "evidence": node["evidence"]})
            flags.append(f"requirement-vs-template: {path} 模板实测 {old}，"
                         f"格式要求为 {new}（{node['evidence'][:40]}）→ 按要求执行")
        _set_rule(merged, path, node)
    return {"profile": merged, "conflicts": conflicts,
            "applied": len(parsed["rules"]), "unparsed": parsed.get("unparsed", [])}


def _differs(a, b) -> bool:
    if isinstance(a, dict) and isinstance(b, dict):
        ka = {k: v for k, v in a.items() if k not in ("source", "evidence", "what")}
        kb = {k: v for k, v in b.items() if k not in ("source", "evidence", "what")}
        return ka != kb
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) > 1e-6  # 要求为精确规格，不做容差吞并
    return a != b


_WRAPPER_KEYS = {"value", "source", "evidence", "what", "confidence"}


def normalize_profile(profile):
    """应用侧入口：把 {"value":..,"source":..} 包装还原为裸值（递归）。
    含非包装键的 dict（如 line_spacing 的 rule/value）原样保留。"""
    if isinstance(profile, dict):
        if set(profile) <= _WRAPPER_KEYS and "value" in profile:
            return profile["value"]
        return {k: normalize_profile(v) for k, v in profile.items()}
    if isinstance(profile, list):
        return [normalize_profile(x) for x in profile]
    return profile


# ---------------------------------------------------------------- CLI
def main():
    if hasattr(sys.stdout, "buffer"):
        sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8",
                                      errors="replace")
    ap = argparse.ArgumentParser()
    ap.add_argument("input", nargs="?", help="要求文件（docx/pdf/txt/md/png/jpg）")
    ap.add_argument("--text", help="直接解析文字")
    ap.add_argument("-o", "--out", help="输出 json 路径")
    ap.add_argument("--merge", help="合并进该 profile.json")
    ap.add_argument("--levels", help="泛指标题时模板在用的级别，如 1,2,3")
    args = ap.parse_args()
    if not args.input and not args.text:
        ap.error("需要文件或 --text")
    text = args.text or extract_text(args.input)
    levels = [int(x) for x in args.levels.split(",")] if args.levels else None
    parsed = RequirementParser(levels).parse(text)
    if args.merge:
        profile = json.load(open(args.merge, encoding="utf-8"))
        result = merge_profile(profile, parsed)
        out = args.out or args.merge
        json.dump(result["profile"], open(out, "w", encoding="utf-8"),
                  ensure_ascii=False, indent=1)
        print(f"已合并 {result['applied']} 条要求规则 -> {out}")
        for c in result["conflicts"]:
            print(f"  ⚠ {c['path']}: 模板 {c['template']} vs 要求 {c['requirement']}")
        for u in result["unparsed"]:
            print(f"  · 未解析：{u[:60]}")
    else:
        out = args.out or (os.path.splitext(args.input or "requirements")[0] + ".req.json")
        json.dump(parsed, open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        print(f"{len(parsed['rules'])} 条规则 / {len(parsed['unparsed'])} 句未解析 -> {out}")
        for path, node in parsed["rules"].items():
            print(f"  {path} = {node['value']}  ← {node['evidence'][:40]}")


if __name__ == "__main__":
    main()
