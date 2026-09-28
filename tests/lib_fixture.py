# -*- coding: utf-8 -*-
"""lib_fixture.py — 夹具构造库

用 python-docx + 原生 OOXML 构造已知格式属性的模板/目标文档，
ground truth 由构造过程决定，评测时按平铺路径比对。
"""
from __future__ import annotations

import os
import struct
import zlib

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

# ---------------------------------------------------------------- 小工具
def make_png(path, w=120, h=80, rgb=(70, 130, 180)):
    def chunk(tag, data):
        c = struct.pack(">I", len(data)) + tag + data
        return c + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)

    ihdr = struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)
    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))
    png = (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) +
           chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))
    with open(path, "wb") as f:
        f.write(png)
    return path


def _rfonts(rpr, east=None, ascii_=None):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    if ascii_:
        rf.set(qn("w:ascii"), ascii_)
        rf.set(qn("w:hAnsi"), ascii_)
    if east:
        rf.set(qn("w:eastAsia"), east)
    return rf


def style_para(doc, text, east=None, ascii_=None, size=None, bold=None,
               align=None, line_spacing=None, first_line_chars=None,
               space_before=None, space_after=None, keep_next=None,
               style=None):
    p = doc.add_paragraph(style=style)
    if text:
        run = p.add_run(text)
        rpr = run._r.get_or_add_rPr()
        if east or ascii_:
            _rfonts(rpr, east, ascii_)
        if size:
            run.font.size = Pt(size)
        if bold is not None:
            run.font.bold = bold
    pf = p.paragraph_format
    if align is not None:
        pf.alignment = align
    if line_spacing is not None:
        pf.line_spacing = line_spacing
    if space_before is not None:
        pf.space_before = Pt(space_before)
    if space_after is not None:
        pf.space_after = Pt(space_after)
    if keep_next:
        ppr = p._p.get_or_add_pPr()
        ppr.append(OxmlElement("w:keepNext"))
    if first_line_chars is not None:
        ind = p._p.get_or_add_pPr().find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            p._p.get_or_add_pPr().append(ind)
        ind.set(qn("w:firstLineChars"), str(int(first_line_chars * 100)))
    return p


def set_pgnum(section, fmt=None, start=None):
    sectPr = section._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sectPr.append(pg)
    if fmt:
        pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))


def footer_page_field(section, instr=r"PAGE \* arabic \* MERGEFORMAT",
                      align=WD_ALIGN_PARAGRAPH.CENTER, east="宋体", size=9,
                      literal="1"):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.paragraph_format.alignment = align
    run = p.add_run()
    run.font.size = Pt(size)
    _rfonts(run._r.get_or_add_rPr(), east)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), f" {instr} ")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = literal
    r.append(t)
    fld.append(r)
    p._p.append(fld)
    return p


def a4(margins=(2.54, 2.54, 3.17, 3.17), section=None):
    sec = section
    if sec is not None:
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin, sec.bottom_margin = Cm(margins[0]), Cm(margins[1])
        sec.left_margin, sec.right_margin = Cm(margins[2]), Cm(margins[3])
    return sec


def add_image(doc, png_path, width_cm, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = align
    p.add_run().add_picture(png_path, width=Cm(width_cm))
    return p


def _border(el_parent, side, val, sz=None, color="000000"):
    b = el_parent.find(qn(f"w:{side}"))
    if b is None:
        b = OxmlElement(f"w:{side}")
        el_parent.append(b)
    b.set(qn("w:val"), val)
    if sz:
        b.set(qn("w:sz"), str(sz))
    b.set(qn("w:color"), color)
    b.set(qn("w:space"), "0")
    return b


def three_line_table(doc, rows, header_bold=True, repeat_header=True,
                     cant_split=True, cell_size=10.5, cell_east="宋体"):
    """rows: list[list[str]]，首行为表头。三线：上下 1.5pt，表头下 0.75pt。"""
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = t._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    tblPr.append(borders)
    _border(borders, "top", "single", 12)
    _border(borders, "bottom", "single", 12)
    for side in ("left", "right", "insideH", "insideV"):
        _border(borders, side, "nil")
    for i, row in enumerate(rows):
        trPr = tbl.tr_lst[i].get_or_add_trPr()
        if cant_split:
            trPr.append(OxmlElement("w:cantSplit"))
        if i == 0 and repeat_header:
            trPr.append(OxmlElement("w:tblHeader"))
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.size = Pt(cell_size)
            if i == 0 and header_bold:
                run.font.bold = True
            _rfonts(run._r.get_or_add_rPr(), east=cell_east)
            if i == 0:
                tcb = cell._tc.get_or_add_tcPr()
                cb = OxmlElement("w:tcBorders")
                tcb.append(cb)
                _border(cb, "bottom", "single", 6)
    return t


def grid_table(doc, rows, cell_size=10.5):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.size = Pt(cell_size)
            if i == 0:
                run.font.bold = True
    return t


def code_box_table(doc, lines, font="Consolas", size=10.5, fill="F2F2F2"):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    tcb = cell._tc.get_or_add_tcPr()
    cb = OxmlElement("w:tcBorders")
    tcb.append(cb)
    for side in ("top", "left", "bottom", "right"):
        _border(cb, side, "single", 4)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tcb.append(shd)
    first = True
    for ln in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(ln)
        run.font.size = Pt(size)
        _rfonts(run._r.get_or_add_rPr(), ascii_=font)
    return t


def omml_paragraph(doc, text, size=12, font="Cambria Math", number=None):
    """插入 OMML 公式段；number 如 "(3-1)" 时加右对齐制表位。"""
    p = doc.add_paragraph()
    om = OxmlElement("m:oMath")
    mr = OxmlElement("m:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), font)
    rf.set(qn("w:hAnsi"), font)
    rpr.append(rf)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    mr.append(rpr)
    mt = OxmlElement("m:t")
    mt.text = text
    mr.append(mt)
    om.append(mr)
    p._p.append(om)
    if number:
        run = p.add_run("\t" + number)
        run.font.size = Pt(size)
        ppr = p._p.get_or_add_pPr()
        tabs = OxmlElement("w:tabs")
        for pos, val in ((int(21 / 2 * 567), "center"), (int(21 * 567) - 800, "right")):
            tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), val)
            tab.set(qn("w:pos"), str(pos))
            tabs.append(tab)
        ppr.append(tabs)
    return p


def field_caption(doc, label, chapter_literal, title, east="宋体", size=10.5,
                  align=WD_ALIGN_PARAGRAPH.CENTER, bold=False, sep=" ",
                  styleref=True):
    """域题注：图 {STYLEREF}-{SEQ} 标题（styleref=False 时为全文连续 SEQ 编号）"""
    p = doc.add_paragraph()
    p.paragraph_format.alignment = align
    def _run(text):
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        rf = OxmlElement("w:rFonts")
        rf.set(qn("w:eastAsia"), east)
        rpr.append(rf)
        szel = OxmlElement("w:sz")
        szel.set(qn("w:val"), str(int(size * 2)))
        rpr.append(szel)
        if bold:
            rpr.append(OxmlElement("w:b"))
        r.append(rpr)
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        r.append(t)
        return r
    def _fld(instr, literal):
        f = OxmlElement("w:fldSimple")
        f.set(qn("w:instr"), f" {instr} ")
        f.append(_run(literal))
        return f
    p._p.append(_run(label))
    if styleref:
        p._p.append(_fld("STYLEREF 1 \\s", chapter_literal))
        p._p.append(_run("-"))
    p._p.append(_fld(f"SEQ {label} \\* ARABIC 1", "1"))
    if sep:
        p._p.append(_run(sep))
    p._p.append(_run(title))
    return p


def manual_caption(doc, text, east="宋体", size=10.5,
                   align=WD_ALIGN_PARAGRAPH.CENTER, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = align
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    _rfonts(run._r.get_or_add_rPr(), east=east)
    return p


def toc_field(doc, depth=3):
    p = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), f' TOC \\o "1-{depth}" \\h \\z \\u ')
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "（目录占位）"
    r.append(t)
    fld.append(r)
    p._p.append(fld)
    return p


def hanging_para(doc, text, size=10.5, hanging_chars=2, east="宋体"):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    _rfonts(run._r.get_or_add_rPr(), east=east)
    ind = p._p.get_or_add_pPr().find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        p._p.get_or_add_pPr().append(ind)
    ind.set(qn("w:hangingChars"), str(int(hanging_chars * 100)))
    return p


def page_break_before(p):
    ppr = p._p.get_or_add_pPr()
    ppr.append(OxmlElement("w:pageBreakBefore"))
    return p


def new_section(doc, start_type=WD_SECTION.NEW_PAGE):
    return doc.add_section(start_type)


def display_formula(doc, text, size=12, number=None):
    """独立行公式：oMath 外再包 oMathPara。"""
    p = omml_paragraph(doc, text, size=size, number=number)
    om = p._p.find(qn("m:oMath"))
    idx = list(p._p).index(om)
    para_el = OxmlElement("m:oMathPara")
    p._p.remove(om)
    para_el.append(om)
    p._p.insert(idx, para_el)
    return p


def merged_three_line(doc, rows, merge_header_cols=(1, 2)):
    """带合并单元格的三线表：表头后两列合并。"""
    t = three_line_table(doc, rows)
    hdr_cells = t.rows[0].cells
    a, b = merge_header_cols
    if b < len(hdr_cells):
        merged = hdr_cells[a].merge(hdr_cells[b])
    return t


def add_anchor_image(doc, png_path, width_cm):
    """浮动图片（wp:anchor，随文锚定）。"""
    from lxml import etree as _et
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_picture(png_path, width=Cm(width_cm))
    drawing = p._p.find(".//" + qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))
    keep = {}
    for ch in list(inline):
        keep[_et.QName(ch).localname] = ch
        inline.remove(ch)
    anchor = OxmlElement("wp:anchor")
    for k, v in dict(distT="0", distB="0", distL="114300", distR="114300",
                     simplePos="0", relativeHeight="2", behindDoc="0",
                     locked="0", layoutInCell="1", allowOverlap="1").items():
        anchor.set(k, v)
    sp = OxmlElement("wp:simplePos")
    sp.set("x", "0")
    sp.set("y", "0")
    anchor.append(sp)
    ph = OxmlElement("wp:positionH")
    ph.set("relativeFrom", "column")
    po = OxmlElement("wp:posOffset")
    po.text = "0"
    ph.append(po)
    anchor.append(ph)
    pv = OxmlElement("wp:positionV")
    pv.set("relativeFrom", "paragraph")
    po2 = OxmlElement("wp:posOffset")
    po2.text = "0"
    pv.append(po2)
    anchor.append(pv)
    for tag in ("extent", "effectExtent"):
        if tag in keep:
            anchor.append(keep[tag])
    anchor.append(OxmlElement("wp:wrapNone"))
    for tag in ("docPr", "graphic"):
        if tag in keep:
            anchor.append(keep[tag])
    drawing.remove(inline)
    drawing.append(anchor)
    return p
