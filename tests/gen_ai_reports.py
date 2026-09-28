# -*- coding: utf-8 -*-
"""gen_ai_reports.py — 基于《数据挖掘实验一》真实内容，按 10 种 AI 文档生成器风格档案
生成实验报告（round1 十份 + round2 十份变体）。

内容种子：实验1_参考答案.md（金融坏账数据预处理与可视化）。
"""
from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lib_fixture import (  # noqa: E402
    Document, WD_ALIGN_PARAGRAPH, Cm, Pt, a4, add_image, code_box_table,
    grid_table, hanging_para, make_png, manual_caption, omml_paragraph,
    style_para, three_line_table, _rfonts, _border,
)
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn as _qn  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "ai_reports")

# ---------------------------------------------------------------- 真实实验内容
PURPOSE = [
    "掌握使用 pandas 进行数据加载、初步探索与清洗的完整流程",
    "掌握检测和处理缺失值、异常值、重复值的常用方法",
    "学会使用 matplotlib 绘制基本统计图表，并进行数据分析",
]
INTRO = ("本实验使用 Kaggle 信贷数据集（银行坏账.csv，共 150000 行 6 列），"
         "围绕数据读取、查询、统计、清洗与可视化五个任务展开，"
         "重点比较正常客户与坏账客户的特征差异。")
CLEAN_TABLE = [["步骤", "剩余行数", "本步删掉"],
               ["原始读取", "150000", "—"],
               ["去重（按客户特征列）", "145900", "4100"],
               ["删月收入缺失", "119762", "26138"],
               ["年龄 20~80", "116675", "3087"],
               ["负债率 ≤1", "109576", "7099"],
               ["月收入 IQR 法", "104881", "4695"]]
NULL_TABLE = [["列", "缺失数", "处理方法"],
              ["月收入", "26138", "删除该行（缺失约 18%，关键列）"],
              ["家属数量", "3380", "中位数填充（缺失仅 2.3%）"]]
STATS_TEXT = ("年龄均值 52.3、中位数 52.0、标准差 14.77，最大 109、最小 0 属异常；"
              "好坏客户计数 139974:10026，坏账占比约 7%，类别不均衡；"
              "分组统计显示坏账客户平均年龄更小（45.9<52.8）、平均月收入更低（5631<6748）。")
FIG_CONCLUSIONS = [
    "图1结论：信用好客户约 93%，坏账约 7%，属类别不均衡数据。",
    "图2结论：坏账率随年龄单调下降，20~33 岁最高约 11.1%，68~80 岁最低约 2.6%。",
    "图3结论：坏账客户月收入中位数 4540 低于好客户 5525，收入越低越易坏账。",
]
CODE_LINES = [
    'df = pd.read_csv("银行坏账.csv", encoding="gbk")',
    'feat_cols = ["年龄", "负债率", "月收入", "家属数量"]',
    "df = df.drop_duplicates(subset=feat_cols, keep=\"first\")",
    'df = df.dropna(subset=["月收入"])',
    'Q1, Q3 = df["月收入"].quantile(0.25), df["月收入"].quantile(0.75)',
    "df = df.loc[(df[\"月收入\"] >= Q1 - 1.5 * (Q3 - Q1)) & (df[\"月收入\"] <= Q3 + 1.5 * (Q3 - Q1))]",
    'counts = df["好坏客户"].value_counts().sort_index()',
    'counts.plot(kind="pie", autopct="%1.1f%%")',
]
CONCLUSION = ("本实验完整走通了数据清洗流程，150000 行数据最终保留 104881 行；"
              "发现坏账客户呈现年轻化、低收入特征，为后续建模提供了方向。")

PNGS = {}


def _png(i, rgb):
    if i not in PNGS:
        PNGS[i] = make_png(os.path.join(OUT, f"_fig{i}.png"), w=140, h=90, rgb=rgb)
    return PNGS[i]


# ---------------------------------------------------------------- 风格档案
def P(**kw):
    base = dict(
        name="style", heading_mode="style", heading_pattern="dotted",
        head_font="黑体", head_sizes=(16, 14, 13), body_font="宋体",
        body_ascii=None, body_size=12, indent_chars=2.0, line_spacing=1.5,
        caption_sep=" ", caption_position="below", caption_numbering="seq",
        caption_font="宋体", caption_size=10.5, table="grid", code="plain",
        code_font="Consolas", image_widths=(13.0, 13.0, 13.0), bold_head=True,
        quirk=None)
    base.update(kw)
    return base


ROUND1 = [
    P(name="docxjs", heading_pattern="dotted", body_size=12, line_spacing=1.3,
      table="plain", code="shade"),
    P(name="pydocx_raw", heading_mode="bold", heading_pattern="dotted",
      body_font="Calibri", body_ascii="Calibri", body_size=11, indent_chars=None,
      line_spacing=1.15, table="grid", code="plain", code_font="Courier New"),
    P(name="pandoc", heading_mode="style", heading_pattern="dotted",
      body_ascii="Calibri", caption_sep=": ", caption_numbering="seq",
      table="plain", code="plain", code_font="Courier New", body_size=12),
    P(name="wps_ai", body_font="楷体", heading_pattern="dotted", head_font="楷体",
      body_size=12, caption_sep=":", caption_numbering="seq", table="grid",
      code="plain", quirk="no_space_after_num"),
    P(name="word_web", body_font="Calibri", body_ascii="Calibri", body_size=11,
      indent_chars=None, line_spacing=1.08, heading_mode="style", table="grid",
      code="shade", code_font="Consolas"),
    P(name="chatgpt", heading_pattern="cn_list", caption_sep="\u3000",
      caption_numbering="seq", table="grid", code="plain", code_font="Courier New"),
    P(name="tongyi", caption_position="above", table="grid", code="boxed",
      body_size=12, heading_mode="style"),
    P(name="latex2docx", body_ascii="Times New Roman", caption_sep=": ",
      table="threeline_thin", code="plain", code_font="Courier New",
      formula=True, body_size=12),
    P(name="md_paste", heading_mode="plain", heading_pattern="none",
      caption_sep="", caption_numbering="seq", caption_font="宋体",
      image_widths=(15.0, 15.0, 15.0), table="grid", code="plain"),
    P(name="manual_ok", heading_mode="style", table="threeline", code="boxed",
      quirk="three_outliers"),
]

ROUND2 = [
    P(name="v2_chapter_manual", heading_pattern="dotted", caption_numbering="chapter",
      caption_sep="：", table="threeline", code="boxed"),
    P(name="v2_auto_num", heading_pattern="dotted", heading_mode="style",
      caption_sep=" ", caption_numbering="chapter", table="grid", code="shade"),
    P(name="v2_english_cap", caption_sep=" ", caption_numbering="seq",
      caption_font="Times New Roman", quirk="english_captions", table="plain",
      code="plain"),
    P(name="v2_caption_above", caption_position="above", caption_sep=" ",
      table="threeline", code="boxed", body_size=10.5),
    P(name="v2_no_caption", caption_numbering="seq", quirk="captions_as_body",
      table="grid", code="plain"),
    P(name="v2_fullwidth_num", heading_pattern="cn_list", caption_sep="：",
      caption_numbering="seq", table="grid", code="boxed"),
    P(name="v2_tight", body_size=10.5, indent_chars=None, line_spacing=1.0,
      caption_size=9, table="plain", code="plain", image_widths=(9.0, 9.0, 9.0)),
    P(name="v2_formula_heavy", formula=True, caption_sep=" ", caption_numbering="chapter",
      table="threeline_thin", code="boxed", code_font="Consolas"),
    P(name="v2_bold_body", heading_mode="bold", heading_pattern="dotted",
      body_font="微软雅黑", body_ascii="微软雅黑", caption_sep=" ", table="grid",
      code="plain"),
    P(name="v2_mixed", heading_pattern="dotted", caption_sep=" ",
      caption_numbering="chapter", quirk="mixed_numbering", table="threeline",
      code="boxed", image_widths=(12.0, 8.0, 12.0)),
]


# ---------------------------------------------------------------- 构造器
def _heading(doc, prof, num, title, level):
    if prof["heading_pattern"] == "dotted" and num != "":
        text = f"{num} {title}"
    elif prof["heading_pattern"] == "cn_list" and isinstance(num, int):
        CN = "一二三四五六七八九十"
        text = f"{CN[num - 1]}、{title}" if level == 1 else f"（{CN[num - 1]}）{title}"
    else:
        text = title
    style = None
    if prof["heading_mode"] == "style":
        style = f"Heading {level}"
    p = style_para(doc, text, east=prof["head_font"] if prof["heading_mode"] != "plain"
                   else None, ascii_=prof["body_ascii"],
                   size=prof["head_sizes"][level - 1],
                   bold=prof["bold_head"] if prof["heading_mode"] != "plain" else False,
                   style=style)
    return p


def _caption(doc, prof, kind, num_str, title, above):
    sep = prof["caption_sep"]
    if prof.get("quirk") == "no_space_after_num":
        sep = ":"
    text = f"{kind}{num_str}{sep}{title}"
    if prof.get("quirk") == "english_captions":
        kind_en = "Figure" if kind == "图" else "Table"
        text = f"{kind_en} {num_str}{sep}{title}"
    p = manual_caption(doc, text, east=prof["caption_font"],
                       size=prof["caption_size"])
    return p


def _body(doc, prof, text):
    kw = dict(east=prof["body_font"], ascii_=prof["body_ascii"],
              size=prof["body_size"], align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    if prof["line_spacing"]:
        kw["line_spacing"] = prof["line_spacing"]
    if prof["indent_chars"]:
        kw["first_line_chars"] = prof["indent_chars"]
    else:
        kw["first_line_chars"] = 0
    return style_para(doc, text, **kw)


def _table(doc, prof, rows, title_num):
    mode = prof["table"]
    if mode == "threeline" or (mode == "threeline_thin"):
        t = three_line_table(doc, rows)
        if mode == "threeline_thin":
            tbl = t._tbl
            borders = tbl.tblPr.find(_qn("w:tblBorders"))
            for side in ("top", "bottom"):
                _border(borders, side, "single", 6)
            for tc in tbl.tr_lst[0].findall(_qn("w:tc")):
                cb = tc.find(_qn("w:tcPr") + "/" + _qn("w:tcBorders"))
                _border(cb, "bottom", "single", 4)
        return t
    if mode == "plain":
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                cell = t.cell(i, j)
                cell.text = ""
                run = cell.paragraphs[0].add_run(v)
                run.font.size = Pt(10.5)
                if i == 0:
                    run.font.bold = True
        return t
    return grid_table(doc, rows)


def _code(doc, prof, lines):
    mode = prof["code"]
    if mode == "boxed":
        return code_box_table(doc, lines[:6], font=prof["code_font"])
    for ln in lines[:6]:
        p = style_para(doc, ln, ascii_=prof["code_font"], east=prof["code_font"],
                       size=10.5, line_spacing=1.0)
        if mode == "shade":
            ppr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(_qn("w:val"), "clear")
            shd.set(_qn("w:fill"), "F5F5F5")
            ppr.append(shd)
    return None


def build_report(prof: dict):
    doc = Document()
    a4(section=doc.sections[0])
    style_para(doc, "数据挖掘与应用课程实验报告——实验一 金融坏账数据预处理和可视化",
               east=prof["head_font"], size=prof["head_sizes"][0] + 2, bold=True,
               align=WD_ALIGN_PARAGRAPH.CENTER)
    _heading(doc, prof, 1, "实验目的", 1)
    for i, pu in enumerate(PURPOSE, 1):
        style_para(doc, f"{i}. {pu}", east=prof["body_font"],
                   ascii_=prof["body_ascii"], size=prof["body_size"],
                   line_spacing=prof["line_spacing"])
    _heading(doc, prof, 2, "实验内容", 1)
    _body(doc, prof, INTRO)
    _heading(doc, prof, 1, "实验步骤", 1) if False else None
    n2 = 1
    _heading(doc, prof, f"3.{n2}" if prof["heading_pattern"] == "dotted" else "",
             "数据读取与清洗", 2)
    n2 += 1
    _code(doc, prof, CODE_LINES[:4])
    _body(doc, prof, "清洗流程与行数对照如下表所示，重复值以四个客户特征列判断，"
                     "共删除 4100 行；月收入缺失占 17.9% 直接删行，家属数量以中位数填充。")
    _table(doc, prof, CLEAN_TABLE, 1)
    if prof.get("quirk") != "captions_as_body":
        _caption(doc, prof, "表", "1", "数据清洗流程对照", above=True)
    else:
        _body(doc, prof, "表1 数据清洗流程对照")
    _heading(doc, prof, f"3.{n2}" if prof["heading_pattern"] == "dotted" else "",
             "数据统计", 2)
    n2 += 1
    _body(doc, prof, STATS_TEXT)
    _table(doc, prof, NULL_TABLE, 2)
    if prof.get("quirk") != "captions_as_body":
        _caption(doc, prof, "表", "2", "缺失值处理方案", above=True)
    else:
        _body(doc, prof, "表2 缺失值处理方案")
    if prof.get("formula"):
        omml_paragraph(doc, "IQR=Q3−Q1, 正常范围=[Q1−1.5IQR, Q3+1.5IQR]", number="(1)")
    _heading(doc, prof, f"3.{n2}" if prof["heading_pattern"] == "dotted" else "",
             "数据可视化", 2)
    widths = prof["image_widths"]
    for i, (rgb, concl) in enumerate(zip(((70, 130, 180), (180, 120, 70), (110, 170, 110)),
                                         FIG_CONCLUSIONS)):
        w = widths[i]
        if prof.get("quirk") == "three_outliers":
            w = (8.0, 13.0, 13.0)[i]
        cap_num = f"1-{i + 1}" if prof["caption_numbering"] == "chapter" else str(i + 1)
        if prof.get("quirk") == "mixed_numbering":
            cap_num = str(i + 1) if i != 1 else f"1-{i + 1}"  # 混用两种编号
        title = concl.split("：", 1)[1][:10]
        if prof.get("quirk") == "three_outliers" and i == 1:
            cap_num = f"1{i + 1}"  # 漏分隔（12 而不是 1-2 / 2）
        if prof["caption_position"] == "above" and prof.get("quirk") != "captions_as_body":
            _caption(doc, prof, "图", cap_num, title, above=True)
        add_image(doc, _png(i, rgb), w)
        if prof["caption_position"] == "below" and prof.get("quirk") != "captions_as_body":
            _caption(doc, prof, "图", cap_num, title, above=False)
        _body(doc, prof, concl)
    _heading(doc, prof, 4, "实验结论", 1)
    _body(doc, prof, CONCLUSION)
    return doc


def main():
    os.makedirs(os.path.join(OUT, "round1"), exist_ok=True)
    os.makedirs(os.path.join(OUT, "round2"), exist_ok=True)
    for prof in ROUND1:
        build_report(prof).save(os.path.join(OUT, "round1", f"{prof['name']}.docx"))
        print("round1", prof["name"])
    for prof in ROUND2:
        build_report(prof).save(os.path.join(OUT, "round2", f"{prof['name']}.docx"))
        print("round2", prof["name"])


if __name__ == "__main__":
    main()
