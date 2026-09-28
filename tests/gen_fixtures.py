# -*- coding: utf-8 -*-
"""gen_fixtures.py — 生成测试夹具（template.docx / target.docx / expected.json）

用法：python gen_fixtures.py [场景名 ...]（缺省全部）
"""
from __future__ import annotations

import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lib_fixture import (  # noqa: E402
    Document, WD_ALIGN_PARAGRAPH, Cm, Pt, a4, add_image, code_box_table,
    field_caption, footer_page_field, grid_table, hanging_para, make_png,
    manual_caption, new_section, omml_paragraph, page_break_before,
    set_pgnum, style_para, three_line_table, toc_field,
)

HERE = os.path.dirname(os.path.abspath(__file__))
FIXDIR = os.path.join(HERE, "fixtures")
SCENARIOS = {}


def scenario(name):
    def deco(fn):
        SCENARIOS[name] = fn
        return fn
    return deco


def write_case(name, template_doc, target_doc, expected):
    d = os.path.join(FIXDIR, name)
    os.makedirs(d, exist_ok=True)
    template_doc.save(os.path.join(d, "template.docx"))
    target_doc.save(os.path.join(d, "target.docx"))
    with open(os.path.join(d, "expected.json"), "w", encoding="utf-8") as f:
        json.dump(expected, f, ensure_ascii=False, indent=1)
    print("built", name)


BODY = "本文针对复杂系统建模问题，提出了基于层次分析的解决方案，并通过仿真实验验证了方法的有效性与稳定性，实验表明该方案在多种工况下均具有良好表现。"


# ================================================================ 场景 1
@scenario("thesis_standard")
def s1():
    png = make_png(os.path.join(FIXDIR, "_img.png"))
    doc = Document()
    a4(section=doc.sections[0])
    # 封面
    style_para(doc, "基于深度学习的系统建模研究", east="黑体", size=22, bold=True,
               align=WD_ALIGN_PARAGRAPH.CENTER)
    # 目录节（罗马页码）
    sec2 = new_section(doc)
    a4(section=sec2)
    set_pgnum(sec2, fmt="lowerRoman")
    footer_page_field(sec2, instr=r"PAGE \* roman \* MERGEFORMAT", literal="i")
    style_para(doc, "目 录", east="黑体", size=16, bold=True,
               align=WD_ALIGN_PARAGRAPH.CENTER)
    toc_field(doc, depth=3)
    # 正文节（阿拉伯页码从 1）
    sec3 = new_section(doc)
    a4(section=sec3)
    set_pgnum(sec3, fmt="decimal", start=1)
    footer_page_field(sec3, instr=r"PAGE \* arabic \* MERGEFORMAT", literal="1")

    for ch in (1, 2, 3):
        style_para(doc, f"第{ch}章 实验研究" if ch > 1 else "第1章 绪论",
                   east="黑体", ascii_="Times New Roman", size=16, bold=True,
                   align=WD_ALIGN_PARAGRAPH.CENTER, space_before=24, space_after=18,
                   line_spacing=1.5, style="Heading 1")
        style_para(doc, f"{ch}.1 研究内容概述", east="黑体", ascii_="Times New Roman",
                   size=14, bold=True, space_before=13, space_after=13,
                   line_spacing=1.5, style="Heading 2")
        style_para(doc, f"{ch}.1.1 具体方法说明", east="黑体", ascii_="Times New Roman",
                   size=12, bold=True, space_before=13, space_after=6,
                   line_spacing=1.5, style="Heading 3")
        for _ in range(2):
            style_para(doc, BODY, east="宋体", ascii_="Times New Roman", size=12,
                       align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5,
                       first_line_chars=2)
        add_image(doc, png, 14.0)
        field_caption(doc, "图", str(ch), "系统总体架构示意图")
        style_para(doc, BODY, east="宋体", ascii_="Times New Roman", size=12,
                   align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5,
                   first_line_chars=2)
        field_caption(doc, "表", str(ch), "实验参数设置", bold=True)
        three_line_table(doc, [["参数", "取值", "说明"],
                               ["学习率", "0.001", "Adam 优化器"],
                               ["批大小", "64", "双向循环网络"]])
        omml_paragraph(doc, "L(θ)=∑logp(y|x;θ)", number=f"({ch}-1)")
        style_para(doc, BODY, east="宋体", ascii_="Times New Roman", size=12,
                   align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5,
                   first_line_chars=2)
    style_para(doc, "参考文献", east="黑体", size=16, bold=True,
               align=WD_ALIGN_PARAGRAPH.CENTER, style="Heading 1")
    for i in (1, 2, 3):
        hanging_para(doc, f"[{i}] 张三, 李四. 深度学习方法研究[J]. 计算机学报, 2024, 47(2): 1-15.")
    code_box_table(doc, ["import torch", "model = Model(depth=6)",
                         "loss = model.fit(data)  # 训练"])

    expected = {
        "page.size.width_cm": 21.0, "page.size.height_cm": 29.7,
        "page.size.orientation": "portrait",
        "page.margins.top_cm": 2.54, "page.margins.left_cm": 3.17,
        "page_numbers.front_matter.format": "lowerRoman",
        "page_numbers.body.format": "decimal",
        "page_numbers.body.position": "footer-center",
        "page_numbers.cover.has_page_number": False,
        "headings.levels_in_use": [1, 2, 3],
        "headings.level1.size_pt.value": 16.0,
        "headings.level1.font_eastasia.value": "黑体",
        "headings.level1.numbering.pattern": "第{n}章",
        "headings.level1.numbering.mode": "manual",
        "headings.level2.size_pt.value": 14.0,
        "headings.level2.numbering.pattern": "{n}.{n2}",
        "headings.level3.size_pt.value": 12.0,
        "body.size_pt.value": 12.0, "body.font_eastasia.value": "宋体",
        "body.first_line_indent_chars.value": 2.0,
        "body.line_spacing.value": 1.5, "body.align.value": "both",
        "figures.width_rule.value": 14.0, "figures.align.value": "center",
        "captions.figure.position": "below", "captions.figure.separator": " ",
        "captions.figure.numbering.mode": "field-seq-styleref",
        "captions.figure.numbering.chapter_linked": True,
        "captions.figure.size_pt.value": 10.5,
        "captions.table.position": "above", "captions.table.separator": " ",
        "captions.table.bold.value": True,
        "tables.three_line.enabled": True,
        "tables.repeat_header.value": True, "tables.cant_split_rows.value": True,
        "tables.continuation_convention.detected": False,
        "code_blocks.box_style": "single-cell-table",
        "code_blocks.font_ascii.value": "Consolas",
        "bibliography.numbering_format": "[{n}]",
        "bibliography.hanging_indent_chars.value": 2.0,
        "toc.mode": "field", "toc.depth": 3,
        "formulas.counts.omml": {"op": ">=", "value": 3},
    }
    write_case("thesis_standard", doc, _messy_target(), expected)


def _messy_target():
    """通用的乱格式目标文档（用于场景 1/2/3 的 target）。"""
    png = os.path.join(FIXDIR, "_img.png")
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for attr, v in (("top_margin", 2.5), ("bottom_margin", 2.5),
                    ("left_margin", 2.5), ("right_margin", 2.5)):
        setattr(sec, attr, Cm(v))
    style_para(doc, "系统建模研究（草稿）", east="楷体", size=18, bold=True,
               align=WD_ALIGN_PARAGRAPH.CENTER)
    for i in (1, 2):
        style_para(doc, f"一、第{i}部分研究内容", east="楷体", size=15, bold=True)
        style_para(doc, BODY, east="楷体", size=12, align=WD_ALIGN_PARAGRAPH.LEFT,
                   line_spacing=1.0)
        manual_caption(doc, f"图 {i}：系统架构示意")      # 题注在图上 + 全角冒号 + 连续编号
        add_image(doc, png, 8.0, align=WD_ALIGN_PARAGRAPH.LEFT)
        manual_caption(doc, f"表 {i}：参数设置")
        grid_table(doc, [["参数", "取值"], ["学习率", "0.001"]])
        style_para(doc, "import numpy as np", ascii_="Consolas", size=12)
        style_para(doc, "model.fit(data)", ascii_="Consolas", size=12)
    style_para(doc, "参考文献", east="楷体", size=15, bold=True)
    style_para(doc, "[1] 张三. 深度学习研究[J]. 计算机学报, 2024.", size=12)
    return doc


# ================================================================ 场景 2
@scenario("mathmodel_manual")
def s2():
    """数模风格：手打连续编号题注、代码在附录、无公式。"""
    png = make_png(os.path.join(FIXDIR, "_img2.png"), rgb=(150, 100, 60))
    doc = Document()
    a4(margins=(2.5, 2.5, 2.5, 3.0), section=doc.sections[0])
    style_para(doc, "摘 要", east="黑体", size=16, bold=True,
               align=WD_ALIGN_PARAGRAPH.CENTER)
    style_para(doc, BODY, east="宋体", size=12, line_spacing=1.5,
               first_line_chars=2, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    style_para(doc, "关键词：模型；仿真；优化", east="宋体", size=12)
    for i in (1, 2, 3):
        style_para(doc, f"{i} 问题重述与分析" if i == 1 else f"{i} 模型建立与求解",
                   east="黑体", size=15, bold=True, style="Heading 1")
        style_para(doc, f"{i}.1 子问题分析", east="黑体", size=14, bold=True,
                   style="Heading 2")
        style_para(doc, BODY, east="宋体", size=12, line_spacing=1.5,
                   first_line_chars=2, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        add_image(doc, png, 12.0)
        manual_caption(doc, f"图{i} 求解流程示意")     # 手打、连续编号、空格分隔
        manual_caption(doc, f"表{i} 灵敏度分析结果")
        three_line_table(doc, [["参数", "+10%", "-10%"],
                               ["目标值", "1.02", "0.98"]])
    style_para(doc, "附录", east="黑体", size=15, bold=True, style="Heading 1")
    code_box_table(doc, ["from scipy.optimize import minimize",
                         "res = minimize(f, x0, method='SLSQP')"])
    expected = {
        "page.margins.top_cm": 2.5, "page.margins.left_cm": 2.5,
        "headings.level1.size_pt.value": 15.0,
        "headings.level1.numbering.pattern": "{n}",
        "headings.level2.size_pt.value": 14.0,
        "body.size_pt.value": 12.0, "body.first_line_indent_chars.value": 2.0,
        "figures.width_rule.value": 12.0,
        "captions.figure.position": "below",
        "captions.figure.separator": " ",
        "captions.figure.numbering.mode": "manual",
        "captions.figure.numbering.chapter_linked": False,
        "captions.table.position": "above",
        "tables.three_line.enabled": True,
        "code_blocks.box_style": "single-cell-table",
        "formulas.counts.omml": 0,
    }
    write_case("mathmodel_manual", doc, _messy_target(), expected)


# ================================================================ 场景 3
@scenario("journal_colon")
def s3():
    """期刊风格：表题在上用全角冒号、图题在下用全角冒号、图 10cm、正文 10.5pt。"""
    png = make_png(os.path.join(FIXDIR, "_img3.png"), rgb=(60, 120, 60))
    doc = Document()
    a4(margins=(2.0, 2.0, 2.0, 2.0), section=doc.sections[0])
    style_para(doc, "A Study on System Modeling", ascii_="Times New Roman",
               size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for i in (1, 2):
        style_para(doc, f"{i} Introduction" if i == 1 else f"{i} Methods",
                   ascii_="Times New Roman", size=12, bold=True, style="Heading 1")
        style_para(doc, BODY, east="宋体", size=10.5, line_spacing=1.25,
                   first_line_chars=2, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        manual_caption(doc, f"表{i}：Parameter settings")
        three_line_table(doc, [["Param", "Value"], ["lr", "0.001"]],
                         cell_size=9)
        add_image(doc, png, 10.0)
        manual_caption(doc, f"图{i}：Pipeline of the framework", size=9)
        style_para(doc, BODY, east="宋体", size=10.5, line_spacing=1.25,
                   first_line_chars=2, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    expected = {
        "page.margins.top_cm": 2.0, "page.margins.left_cm": 2.0,
        "body.size_pt.value": 10.5,
        "figures.width_rule.value": 10.0,
        "captions.figure.position": "below", "captions.figure.separator": "：",
        "captions.table.position": "above", "captions.table.separator": "：",
        "captions.figure.size_pt.value": 9.0,
        "tables.three_line.enabled": True,
    }
    write_case("journal_colon", doc, _messy_target(), expected)


# ================================================================ 场景 4：反向（模板手打、目标域编号）
@scenario("reverse_field_manual")
def s4():
    png = make_png(os.path.join(FIXDIR, "_img4.png"), rgb=(90, 90, 160))
    doc = Document()
    a4(section=doc.sections[0])
    for ch in (1, 2):
        style_para(doc, f"{ch} 研究内容与方案", east="黑体", size=15, bold=True,
                   style="Heading 1")
        style_para(doc, BODY, east="宋体", size=12, line_spacing=1.5,
                   first_line_chars=2, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        add_image(doc, png, 12.5)
        manual_caption(doc, f"图{ch * 2 - 1} 系统架构")
        style_para(doc, BODY, east="宋体", size=12, line_spacing=1.5,
                   first_line_chars=2, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        manual_caption(doc, f"表{ch * 2 - 1} 参数设置")
        three_line_table(doc, [["参数", "值"], ["lr", "0.01"]])
    expected = {
        "headings.level1.numbering.pattern": "{n}",
        "figures.width_rule.value": 12.5,
        "captions.figure.numbering.mode": "manual",
        "captions.figure.numbering.chapter_linked": False,
        "captions.figure.separator": " ",
        "tables.three_line.enabled": True,
    }
    # 目标：域编号题注
    tgt = Document()
    a4(section=tgt.sections[0])
    for ch in (1, 2):
        style_para(tgt, f"一、第{ch}部分", east="楷体", size=15, bold=True)
        style_para(tgt, BODY, east="楷体", size=12)
        add_image(tgt, png, 9.0)
        field_caption(tgt, "图", str(ch), "系统架构")
    write_case("reverse_field_manual", doc, tgt, expected)


# ================================================================ 场景 5：目标标题无样式无编号（歧义确认流）
@scenario("unstyled_target")
def s5():
    png = make_png(os.path.join(FIXDIR, "_img5.png"))
    doc = Document()
    a4(section=doc.sections[0])
    for ch in (1, 2):
        style_para(doc, f"第{ch}章 研究" if ch > 1 else "第1章 绪论", east="黑体",
                   size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                   style="Heading 1")
        style_para(doc, BODY, east="宋体", size=12, line_spacing=1.5,
                   first_line_chars=2, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        add_image(doc, png, 14.0)
        manual_caption(doc, f"图{ch}-1 架构图")
    tgt = Document()
    a4(section=tgt.sections[0])
    style_para(tgt, "系统研究（草稿）", east="楷体", size=18, bold=True,
               align=WD_ALIGN_PARAGRAPH.CENTER)
    style_para(tgt, "绪论", east="楷体", size=15, bold=True)      # idx 1
    style_para(tgt, BODY, east="楷体", size=12)
    add_image(tgt, png, 8.0)
    manual_caption(tgt, "图 1：架构")
    style_para(tgt, "相关工作", east="楷体", size=15, bold=True)  # idx 5
    style_para(tgt, BODY, east="楷体", size=12)
    expected = {
        "headings.level1.numbering.pattern": "第{n}章",
        "captions.figure.numbering.chapter_linked": True,
        "captions.figure.separator": " ",
    }
    write_case("unstyled_target", doc, tgt, expected)
    with open(os.path.join(FIXDIR, "unstyled_target", "map_override.json"),
              "w", encoding="utf-8") as f:
        json.dump({"headings": [{"idx": 1, "level": 1, "text": "绪论"},
                                {"idx": 5, "level": 1, "text": "相关工作"}]},
                  f, ensure_ascii=False)


# ================================================================ 场景 6：outlineLvl 识别（WPS 风格）
@scenario("outline_lvl_headings")
def s6():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn as _qn
    png = make_png(os.path.join(FIXDIR, "_img6.png"))
    doc = Document()
    a4(section=doc.sections[0])
    for ch in (1, 2):
        p = style_para(doc, f"{ch} 引言与研究" if ch == 1 else f"{ch} 实验分析",
                       east="黑体", size=15, bold=True)
        ppr = p._p.get_or_add_pPr()
        ol = OxmlElement("w:outlineLvl")
        ol.set(_qn("w:val"), "0")
        ppr.append(ol)
        style_para(doc, BODY, east="宋体", size=12, line_spacing=1.5,
                   first_line_chars=2)
        add_image(doc, png, 11.0)
        manual_caption(doc, f"图{ch} 示意图")
    expected = {
        "headings.level1.numbering.pattern": "{n}",
        "headings.levels_in_use": [1],
        "figures.width_rule.value": 11.0,
    }
    write_case("outline_lvl_headings", doc, _messy_target(), expected)


# ================================================================ 场景 7：横页混排
@scenario("landscape_mixed")
def s7():
    from docx.enum.section import WD_ORIENT
    doc = Document()
    a4(margins=(2.0, 2.0, 2.0, 2.0), section=doc.sections[0])
    style_para(doc, "1 实验", east="黑体", size=15, bold=True, style="Heading 1")
    style_para(doc, BODY, east="宋体", size=12)
    sec2 = new_section(doc)
    sec2.orientation = WD_ORIENT.LANDSCAPE
    sec2.page_width, sec2.page_height = Cm(29.7), Cm(21)
    manual_caption(doc, "表1 大规模实验结果")
    three_line_table(doc, [["指标"] + [f"方法{i}" for i in range(8)],
                           ["Acc"] + [str(0.8 + i * 0.01) for i in range(8)]])
    expected = {
        "page.size.mixed_orientation": True,
        "page.margins.top_cm": 2.0,
        "tables.three_line.enabled": True,
    }
    write_case("landscape_mixed", doc, _messy_target(), expected)


# ================================================================ 场景 8：大量图表（性能+重编号）
@scenario("many_figures")
def s8():
    png = make_png(os.path.join(FIXDIR, "_img8.png"))
    doc = Document()
    a4(section=doc.sections[0])
    for ch in (1, 2, 3):
        style_para(doc, f"第{ch}章 实验", east="黑体", size=16, bold=True,
                   style="Heading 1")
        for i in range(1, 9):
            style_para(doc, BODY, east="宋体", size=12)
            add_image(doc, png, 12.0)
            manual_caption(doc, f"图{ch}-{i} 结果示意{i}")
    for i in range(1, 7):
        manual_caption(doc, f"表1-{i} 汇总{i}")
        three_line_table(doc, [["a", "b"], ["1", "2"]])
    expected = {
        "figures.width_rule.value": 12.0,
        "meta.doc_stats.images": {"op": ">=", "value": 20},
        "meta.doc_stats.captions_figure": {"op": ">=", "value": 20},
        "captions.figure.numbering.chapter_linked": True,
    }
    write_case("many_figures", doc, _messy_target(), expected)


# ================================================================ 场景 9：模板自身不一致（9:1）
@scenario("template_inconsistent")
def s9():
    png = make_png(os.path.join(FIXDIR, "_img9.png"))
    doc = Document()
    a4(section=doc.sections[0])
    for i in range(1, 11):
        add_image(doc, png, 13.0)
        if i == 10:
            manual_caption(doc, f"图{i}离群无空格题注")   # 无空格的离群样本
        else:
            manual_caption(doc, f"图{i} 正常题注")
    expected = {
        "captions.figure.separator": " ",
        "figures.width_rule.value": 13.0,
    }
    write_case("template_inconsistent", doc, _messy_target(), expected)


# ================================================================ 场景 10：合并单元格三线表
@scenario("merged_cells")
def s10():
    from lib_fixture import merged_three_line
    png = make_png(os.path.join(FIXDIR, "_img10.png"))
    doc = Document()
    a4(section=doc.sections[0])
    style_para(doc, "1 实验部分", east="黑体", size=15, bold=True, style="Heading 1")
    manual_caption(doc, "表1 多方法对比")
    merged_three_line(doc, [["指标", "本文方法", "对比方法"],
                            ["精确率", "0.91", "0.85"]])
    style_para(doc, BODY, east="宋体", size=12)
    expected = {"tables.three_line.enabled": True}
    tgt = Document()
    a4(section=tgt.sections[0])
    t = grid_table(tgt, [["指标", "A", "B"], ["P", "0.9", "0.8"]])
    t.cell(0, 1).merge(t.cell(0, 2))
    write_case("merged_cells", doc, tgt, expected)


# ================================================================ 场景 11：代码仅在附录
@scenario("code_appendix_only")
def s11():
    doc = Document()
    a4(section=doc.sections[0])
    style_para(doc, "1 方法", east="黑体", size=15, bold=True, style="Heading 1")
    style_para(doc, BODY, east="宋体", size=12, first_line_chars=2)
    style_para(doc, "附录", east="黑体", size=15, bold=True, style="Heading 1")
    code_box_table(doc, ["import numpy as np", "print(np.__version__)"])
    code_box_table(doc, ["def solve(x):", "    return x ** 2"])
    expected = {
        "code_blocks.box_style": "single-cell-table",
        "code_blocks.font_ascii.value": "Consolas",
    }
    tgt = Document()
    a4(section=tgt.sections[0])
    style_para(tgt, "一、方法", east="楷体", size=15, bold=True)
    style_para(tgt, "import numpy as np", ascii_="Consolas", size=12)
    style_para(tgt, "print(1)", ascii_="Consolas", size=12)
    write_case("code_appendix_only", doc, tgt, expected)


# ================================================================ 场景 12：独立行公式（oMathPara）
@scenario("formula_display")
def s12():
    from lib_fixture import display_formula
    doc = Document()
    a4(section=doc.sections[0])
    style_para(doc, "2 模型推导", east="黑体", size=15, bold=True, style="Heading 1")
    for i in (1, 2, 3):
        display_formula(doc, f"E_{i}=mc^2", number=f"(2-{i})")
        style_para(doc, BODY, east="宋体", size=12)
    expected = {
        "formulas.counts.omathpara": {"op": ">=", "value": 3},
        "formulas.omml_font.value": "Cambria Math",
    }
    tgt = Document()
    a4(section=tgt.sections[0])
    style_para(tgt, "一、推导", east="楷体", size=15, bold=True)
    omml_paragraph(tgt, "E=mc^2", number="(1)")
    write_case("formula_display", doc, tgt, expected)


# ================================================================ 场景 13：浮动图片
@scenario("floating_anchor")
def s13():
    from lib_fixture import add_anchor_image
    png = make_png(os.path.join(FIXDIR, "_img13.png"))
    doc = Document()
    a4(section=doc.sections[0])
    style_para(doc, "1 综述", east="黑体", size=15, bold=True, style="Heading 1")
    style_para(doc, BODY, east="宋体", size=12)
    add_image(doc, png, 13.0)
    manual_caption(doc, "图1 总览")
    expected = {"figures.width_rule.value": 13.0}
    tgt = Document()
    a4(section=tgt.sections[0])
    style_para(tgt, "一、综述", east="楷体", size=15, bold=True)
    style_para(tgt, BODY, east="楷体", size=12)
    add_anchor_image(tgt, png, 9.0)
    manual_caption(tgt, "图 1：总览")
    write_case("floating_anchor", doc, tgt, expected)


# ================================================================ 场景 14：目标层级多于模板
@scenario("target_extra_level")
def s14():
    doc = Document()
    a4(section=doc.sections[0])
    for ch in (1, 2):
        style_para(doc, f"{ch} 章节", east="黑体", size=15, bold=True,
                   style="Heading 1")
        style_para(doc, f"{ch}.1 小节", east="黑体", size=14, bold=True,
                   style="Heading 2")
        style_para(doc, BODY, east="宋体", size=12)
    expected = {
        "headings.levels_in_use": [1, 2],
        "headings.level2.numbering.pattern": "{n}.{n2}",
    }
    tgt = Document()
    a4(section=tgt.sections[0])
    for ch in (1, 2):
        style_para(tgt, f"{ch} 章节", east="楷体", size=15, bold=True)
        style_para(tgt, f"{ch}.1 小节", east="楷体", size=14, bold=True)
        style_para(tgt, f"{ch}.1.1 条目", east="楷体", size=13, bold=True)
        style_para(tgt, BODY, east="楷体", size=12)
    write_case("target_extra_level", doc, tgt, expected)


# ================================================================ 场景 15：交叉引用密集
@scenario("xref_dense")
def s15():
    png = make_png(os.path.join(FIXDIR, "_img15.png"))
    doc = Document()
    a4(section=doc.sections[0])
    for ch in (1, 2):
        style_para(doc, f"第{ch}章 分析", east="黑体", size=16, bold=True,
                   style="Heading 1")
        for i in (1, 2):
            add_image(doc, png, 12.0)
            manual_caption(doc, f"图{ch}-{i} 结果图")
        for i in (1, 2):
            style_para(doc, f"实验结果表明（如图{ch}-{i}所示），该方法显著优于基线，"
                            f"且如图{ch}-{i}所见趋势一致。", east="宋体", size=12,
                       first_line_chars=2)
    expected = {
        "captions.figure.numbering.chapter_linked": True,
        "cross_references.count": {"op": ">=", "value": 4},
    }
    tgt = Document()
    a4(section=tgt.sections[0])
    for ch in (1, 2):
        style_para(tgt, f"一、第{ch}部分", east="楷体", size=15, bold=True)
        add_image(tgt, png, 9.0)
        manual_caption(tgt, f"图 {ch}：结果")
        style_para(tgt, f"如图{ch}所示效果良好，另见图{ch}。", east="楷体", size=12)
    write_case("xref_dense", doc, tgt, expected)


# ================================================================ 场景 16：模板无公式（目标有）
@scenario("no_formulas_template")
def s16():
    png = make_png(os.path.join(FIXDIR, "_img16.png"))
    doc = Document()
    a4(section=doc.sections[0])
    for ch in (1, 2):
        style_para(doc, f"第{ch}章 描述", east="黑体", size=16, bold=True,
                   style="Heading 1")
        style_para(doc, BODY, east="宋体", size=12, first_line_chars=2)
        add_image(doc, png, 12.0)
        manual_caption(doc, f"图{ch}-1 示意")
        manual_caption(doc, f"表{ch}-1 数据")
        three_line_table(doc, [["a", "b"], ["1", "2"]])
    expected = {
        "dimensions_absent": {"op": "in", "value": "formulas"},
        "figures.width_rule.value": 12.0,
    }
    tgt = Document()
    a4(section=tgt.sections[0])
    style_para(tgt, "一、描述", east="楷体", size=15, bold=True)
    style_para(tgt, BODY, east="楷体", size=12)
    omml_paragraph(tgt, "a^2+b^2=c^2", number="(1)")
    add_image(tgt, png, 9.0)
    write_case("no_formulas_template", doc, tgt, expected)


# ================================================================ 场景 17：续表惯例
@scenario("continuation_table")
def s17():
    doc = Document()
    a4(section=doc.sections[0])
    style_para(doc, "3 实验数据", east="黑体", size=15, bold=True, style="Heading 1")
    manual_caption(doc, "表3-1 全量实验记录")
    three_line_table(doc, [["编号", "数值"]] + [[str(i), str(i * 3)] for i in range(30)])
    manual_caption(doc, "续表3-1 全量实验记录（续）")
    three_line_table(doc, [["编号", "数值"]] + [[str(i), str(i * 3)] for i in range(30, 60)])
    expected = {
        "tables.continuation_convention.detected": True,
        "tables.three_line.enabled": True,
    }
    write_case("continuation_table", doc, _messy_target(), expected)


# ================================================================ 场景 18：中文序号标题（一、/（一））
@scenario("chinese_numbering")
def s18():
    png = make_png(os.path.join(FIXDIR, "_img18.png"))
    doc = Document()
    a4(section=doc.sections[0])
    for cn in ("一、", "二、"):
        style_para(doc, f"{cn}研究内容", east="黑体", size=15, bold=True)
        style_para(doc, "（一）细分方向", east="黑体", size=14, bold=True)
        style_para(doc, BODY, east="宋体", size=12, first_line_chars=2)
        add_image(doc, png, 12.0)
        manual_caption(doc, f"图{1 if cn == '一、' else 2} 流程")
    expected = {
        "headings.level2.numbering.pattern": "{c}、",
        "headings.level3.numbering.pattern": "（{c}）",
        "captions.figure.numbering.chapter_linked": False,
        "figures.width_rule.value": 12.0,
    }
    tgt = Document()
    a4(section=tgt.sections[0])
    for ch in (1, 2):
        style_para(tgt, f"一、第{ch}部分", east="楷体", size=15, bold=True)
        style_para(tgt, BODY, east="楷体", size=12)
        add_image(tgt, png, 8.0)
        manual_caption(tgt, f"图 {ch}：流程")
    write_case("chinese_numbering", doc, tgt, expected)


# ================================================================ 场景 19：装订线与非对称边距
@scenario("gutter_margins")
def s19():
    png = make_png(os.path.join(FIXDIR, "_img19.png"))
    doc = Document()
    sec = doc.sections[0]
    a4(margins=(3.0, 3.0, 3.5, 2.5), section=sec)
    style_para(doc, "1 内容", east="黑体", size=15, bold=True, style="Heading 1")
    style_para(doc, BODY, east="宋体", size=12)
    add_image(doc, png, 11.0)
    expected = {
        "page.margins.top_cm": 3.0, "page.margins.left_cm": 3.5,
        "page.margins.right_cm": 2.5,
    }
    write_case("gutter_margins", doc, _messy_target(), expected)


# ================================================================ 场景 20：SEQ 连续编号域题注（无 STYLEREF）
@scenario("seq_continuous_fields")
def s20():
    png = make_png(os.path.join(FIXDIR, "_img20.png"))
    doc = Document()
    a4(section=doc.sections[0])
    for ch in (1, 2):
        style_para(doc, f"{ch} 内容", east="黑体", size=15, bold=True,
                   style="Heading 1")
        add_image(doc, png, 12.0)
        field_caption(doc, "图", None, "连续编号题注", styleref=False)
        style_para(doc, BODY, east="宋体", size=12, first_line_chars=2)
    expected = {
        "captions.figure.numbering.mode": "field-seq",
        "captions.figure.numbering.chapter_linked": False,
        "figures.width_rule.value": 12.0,
    }
    write_case("seq_continuous_fields", doc, _messy_target(), expected)


if __name__ == "__main__":
    names = sys.argv[1:] or list(SCENARIOS)
    for n in names:
        SCENARIOS[n]()