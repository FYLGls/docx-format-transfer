# -*- coding: utf-8 -*-
"""test_req_parse.py — 格式要求解析/合并/OCR/网页链路 测试"""
from __future__ import annotations

import io
import json
import os
import subprocess
import sys
import urllib.request

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SCRIPTS = os.path.join(ROOT, "skill", "scripts")
FIX = os.path.join(HERE, "fixtures")
sys.path.insert(0, SCRIPTS)

from parse_requirements import (RequirementParser, extract_text,  # noqa: E402
                                merge_profile, normalize_profile)

SAMPLE = ("论文一律采用A4纸，页边距上下2.5厘米，左右3厘米。"
          "正文采用宋体小四，1.5倍行距，首行缩进2字符。"
          "一级标题用黑体三号居中，二级标题黑体四号。"
          "图注在图下方，表注在表上方，按图2-1格式编号。"
          "表格采用三线表，跨页时重复表头。页码居中，阿拉伯数字。"
          "目录为三级目录。代码用等宽字体装框。参考文献按GB/T 7714顺序编码。每章另起一页。")
fails = []


def check(name, cond, detail=""):
    print(f"[{'PASS' if cond else 'FAIL'}] {name}" + (f"（{detail}）" if detail and not cond else ""))
    if not cond:
        fails.append(name)


def v(profile, path):
    cur = profile
    for p in path.split("."):
        cur = (cur or {}).get(p)
    if isinstance(cur, dict) and "value" in cur and "rule" not in cur:
        return cur["value"]
    return cur


def vr(flat_rules, path):
    """扁平规则 dict（parse 输出）取值"""
    node = flat_rules.get(path)
    return node.get("value") if node else None


# ---------------- 1) 规则族解析 ----------------
r = RequirementParser([1, 2, 3]).parse(SAMPLE)
rules = r["rules"]
expect = {
    "page.size.width_cm": 21.0, "page.margins.top_cm": 2.5, "page.margins.left_cm": 3.0,
    "body.font_eastasia": "宋体", "body.size_pt": 12.0,
    "body.first_line_indent_chars": 2.0,
    "headings.level1.font_eastasia": "黑体", "headings.level1.size_pt": 16.0,
    "headings.level1.align": "center", "headings.level2.size_pt": 14.0,
    "captions.figure.position": "below", "captions.table.position": "above",
    "captions.figure.numbering.chapter_linked": True,
    "tables.border_rule": "three-line", "tables.repeat_header": True,
    "page_numbers.body.position": "footer-center",
    "page_numbers.body.format": "decimal",
    "toc.depth": 3, "code_blocks.font_ascii": "Consolas",
    "code_blocks.box_style": "single-cell-table",
    "bibliography.numbering_format": "[{n}]",
    "headings.level1.page_break_before": True,
}
for path, exp in expect.items():
    check(f"解析 {path} = {exp}", vr(rules, path) == exp, f"实际 {vr(rules, path)}")
check("行距 1.5 倍", (vr(rules, "body.line_spacing") or {}).get("value") == 1.5)
check("无未解析句", len(r["unparsed"]) == 0, str(r["unparsed"])[:100])
check("每条规则带 evidence", all("evidence" in n for n in rules.values()))

# ---------------- 2) 合并优先级 + 冲突标记 ----------------
tpl_profile = {"page": {"margins": {"top_cm": {"value": 2.54, "confidence": "high"}}},
               "body": {"font_eastasia": {"value": "楷体"},
                        "line_spacing": {"rule": "auto", "value": 1.0}},
               "flags": []}
req = RequirementParser([1, 2, 3]).parse("页边距上下均为2.5厘米。正文宋体，1.5倍行距。")
merged = merge_profile(tpl_profile, req)
mp = merged["profile"]
check("要求覆盖模板 2.54→2.5", v(mp, "page.margins.top_cm") == 2.5)
check("要求覆盖字体 楷体→宋体", v(mp, "body.font_eastasia") == "宋体")
check("覆盖后保留 source=格式要求",
      mp["page"]["margins"]["top_cm"].get("source") == "格式要求")
check("冲突计数 3（边距/字体/行距）", len(merged["conflicts"]) == 3,
      str(merged["conflicts"])[:150])
check("冲突写入 flags", any("requirement-vs-template" in f for f in mp["flags"]))
check("一致项不记冲突（缩进缺失不冲突）",
      not any(c["path"] == "body.first_line_indent_chars" for c in merged["conflicts"]))

# ---------------- 3) normalize（应用侧入口） ----------------
nm = normalize_profile(json.loads(json.dumps(mp)))
check("normalize 解包 wrapper", nm["page"]["margins"]["top_cm"] == 2.5)
check("normalize 保留 line_spacing 结构",
      nm["body"]["line_spacing"].get("value") == 1.5 and
      nm["body"]["line_spacing"].get("rule") == "auto")

# ---------------- 4) docx 提取 + OCR 链路 ----------------
os.makedirs(os.path.join(FIX, "requirements"), exist_ok=True)
req_docx = os.path.join(FIX, "requirements", "格式要求.docx")
from docx import Document  # noqa: E402
d = Document()
d.add_paragraph("实验报告格式要求")
d.add_paragraph(SAMPLE)
d.save(req_docx)
text = extract_text(req_docx)
check("docx 文本提取", "三线表" in text and "页边距" in text)

img_path = os.path.join(FIX, "requirements", "格式要求.png")
try:
    from PIL import Image, ImageDraw, ImageFont
    font = None
    for fp in (r"C:\Windows\Fonts\simhei.ttf", r"C:\Windows\Fonts\msyh.ttc",
               r"C:\Windows\Fonts\simsun.ttc"):
        if os.path.isfile(fp):
            font = ImageFont.truetype(fp, 28)
            break
    lines = SAMPLE.split("。")
    img = Image.new("RGB", (1100, 60 * len(lines) + 40), "white")
    dr = ImageDraw.Draw(img)
    for i, ln in enumerate(lines):
        dr.text((20, 20 + i * 60), ln + "。", fill="black", font=font)
    img.save(img_path)
    ocr_text = extract_text(img_path)
    ocr_rules = RequirementParser([1, 2, 3]).parse(ocr_text)
    hit = vr(ocr_rules["rules"], "page.margins.top_cm") == 2.5 or \
        vr(ocr_rules["rules"], "tables.border_rule") == "three-line"
    check("OCR 图片→解析（rapidocr）", hit,
          f"OCR 文本 {len(ocr_text)} 字，规则 {len(ocr_rules['rules'])} 条")
except Exception as e:
    print(f"[SKIP] OCR 链路：{type(e).__name__}: {e}")

# ---------------- 5) 网页链路 2 轮（detect 带要求 → apply → verify） ----------------
PORT = 8977
BASE = f"http://127.0.0.1:{PORT}"
proc = subprocess.Popen([sys.executable, os.path.join(SCRIPTS, "serve_ui.py"), str(PORT)],
                        stdout=subprocess.DEVNULL, stderr=subprocess.STDOUT)
try:
    for _ in range(60):
        import time
        time.sleep(0.25)
        try:
            urllib.request.urlopen(BASE + "/api/version", timeout=2).read()
            break
        except Exception:
            pass
    else:
        raise SystemExit("服务器未启动")

    def post(path, data):
        req = urllib.request.Request(
            BASE + path, data=json.dumps(data).encode(),
            headers={"Content-Type": "application/json"}, method="POST")
        return urllib.request.urlopen(req, timeout=120)

    FIXT = os.path.join(FIX, "thesis_standard")
    rounds = [
        ("要求覆盖边距", {"page.margins.top_cm": 2.5}),
        ("要求三线表+页码居中", {"tables.border_rule": "three-line"}),
    ]
    for i, (label, key_paths) in enumerate(rounds, 1):
        req_text = SAMPLE if i == 1 else "表格采用三线表。页码居中。"
        steps, req_ev, done = 0, None, None
        for line in post("/api/detect", {
                "path": os.path.join(FIXT, "template.docx"),
                "requirement_text": req_text}).read().decode("utf-8").splitlines():
            if not line.strip():
                continue
            ev = json.loads(line)
            if ev["type"] == "step":
                steps += 1
            elif ev["type"] == "requirement":
                req_ev = ev
            elif ev["type"] == "done":
                done = ev
        errs = []
        if steps != 13:
            errs.append(f"检测步数 {steps}≠13")
        if req_ev is None or req_ev.get("applied", 0) < 3:
            errs.append(f"要求规则过少 {req_ev and req_ev.get('applied')}")
        if req_ev and req_ev.get("error"):
            errs.append("requirement: " + req_ev["error"])
        # apply
        saved = verify = None
        for line in post("/api/apply", {
                "target": os.path.join(FIXT, "target.docx"),
                "profile": done["profile"]}).read().decode("utf-8").splitlines():
            if not line.strip():
                continue
            ev = json.loads(line)
            if ev["type"] == "saved":
                saved = ev
            elif ev["type"] == "verify":
                verify = ev
        if verify is None or verify["fail"] != 0:
            errs.append(f"verify 失败 {verify and verify['fail']}")
        # 要求值在输出中实测生效
        from docx import Document
        from docx.shared import Cm
        doc = Document(saved["output"])
        if i == 1:
            top = doc.sections[0].top_margin.cm
            if abs(top - 2.5) > 0.06:
                errs.append(f"要求边距未生效：{top}")
        else:
            from docx.oxml.ns import qn
            tbl = doc.tables[0]._tbl
            borders = tbl.tblPr.find(qn("w:tblBorders"))
            inside_v = borders.find(qn("w:insideV")) if borders is not None else None
            if inside_v is None or inside_v.get(qn("w:val")) != "nil":
                errs.append("三线表未生效（insideV 非 nil）")
        check(f"网页链路 轮{i} {label}", not errs, "; ".join(errs))
        if errs:
            fails.append(f"web{i}")
finally:
    proc.terminate()

print(f"\n{'全部通过' if not fails else '失败：' + ', '.join(fails)}")
sys.exit(1 if fails else 0)
